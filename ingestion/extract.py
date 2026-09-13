"""Text extraction for supported file types."""

import io
from pathlib import Path

import docx  # python-docx
import pymupdf  # PyMuPDF (formerly imported as `fitz`)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_bytes(path.read_bytes(), ".pdf")
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"Unsupported file type: {suffix}")


def extract_text_from_bytes(data: bytes, suffix: str) -> str:
    suffix = suffix.lower()
    if suffix == ".pdf":
        text_parts = []
        with pymupdf.open(stream=data, filetype="pdf") as pdf:
            for page in pdf:
                text_parts.append(page.get_text())
        return "\n".join(text_parts)
    if suffix == ".docx":
        document = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in document.paragraphs)
    if suffix in (".txt", ".md"):
        return data.decode("utf-8", errors="ignore")
    raise ValueError(f"Unsupported byte extraction for: {suffix}")


def _extract_docx(path: Path) -> str:
    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs)
