"""Unit tests for text extraction, including in-memory DOCX and PDF byte extraction."""

import io
import sys
from pathlib import Path

import docx
import pymupdf
import pytest

sys.path.append(str(Path(__file__).resolve().parent.parent))

from ingestion.extract import extract_text, extract_text_from_bytes


def test_extract_text_from_docx_bytes():
    # Build an in-memory docx
    doc = docx.Document()
    doc.add_heading("Company Leave Policy", level=1)
    doc.add_paragraph("Employees receive 20 days of paid time off annually.")
    doc.add_paragraph("Manager approval is required for leave exceeding 5 days.")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    extracted = extract_text_from_bytes(docx_bytes, ".docx")
    assert "Company Leave Policy" in extracted
    assert "20 days of paid time off" in extracted
    assert "Manager approval" in extracted


def test_extract_text_from_pdf_bytes():
    # Build an in-memory pdf
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Internal Security Policy: Passwords must be 16+ chars.")

    pdf_bytes = doc.tobytes()
    doc.close()

    extracted = extract_text_from_bytes(pdf_bytes, ".pdf")
    assert "Internal Security Policy" in extracted
    assert "16+ chars" in extracted


def test_extract_text_from_txt_bytes():
    raw = b"Sample plain text content from export."
    extracted = extract_text_from_bytes(raw, ".txt")
    assert extracted == "Sample plain text content from export."


def test_extract_text_unsupported_suffix():
    with pytest.raises(ValueError, match="Unsupported byte extraction"):
        extract_text_from_bytes(b"data", ".xlsx")


def test_extract_text_local_txt_file(tmp_path):
    f = tmp_path / "test.txt"
    f.write_text("Local sample policy document.", encoding="utf-8")
    assert extract_text(f) == "Local sample policy document."


def test_extract_text_local_docx_file(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("Testing local docx extraction.")
    docx_file = tmp_path / "test.docx"
    doc.save(str(docx_file))

    assert "Testing local docx extraction." in extract_text(docx_file)
